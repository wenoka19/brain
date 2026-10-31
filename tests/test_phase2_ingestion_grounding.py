"""Phase 2 — gatekeeper, grounded extraction, contradiction linker, engine."""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest

from app.core import engine as E
from app.db import control as control_db
from app.db import tenant as tenant_db
from app.ingestion import fetchers, gatekeeper
from app.ingestion.extractor import chunk_content, extract_grounded_claims
from app.ingestion.gatekeeper import GatekeeperError
from app.ingestion.grounder import ground_quote
from app.ingestion.linker import link_claims
from app.ingestion.providers import (
    DeterministicFetcher,
    FetchedContent,
    HashingEmbedder,
    HeuristicClaimExtractor,
    HeuristicLinkClassifier,
    ProviderError,
)
from app.security.sanitizer import UNTRUSTED_END, wrap_untrusted
from app.security.ssrf import SSRFBlocked

YOUTUBE_URL = "https://www.youtube.com/watch?v=dQw4w9WgXcQ"
REPORT_URL = "https://example.com/report"
SKILL_URL = "https://raw.githubusercontent.com/acme/playbooks/main/SKILL.md"
MALICIOUS_URL = "https://example.com/malicious"


@pytest.fixture
def ctx(cfg, user):
    return E.build_context(user["user_id"])


# ===========================================================================
# gatekeeper — URLs
# ===========================================================================
class TestURLGatekeeper:
    @pytest.mark.parametrize(
        "raw,canonical",
        [
            ("https://youtu.be/dQw4w9WgXcQ?si=abc123", YOUTUBE_URL),
            ("https://m.youtube.com/watch?v=dQw4w9WgXcQ&t=30s", YOUTUBE_URL),
            ("https://www.youtube.com/shorts/dQw4w9WgXcQ", YOUTUBE_URL),
            ("https://www.youtube.com/embed/dQw4w9WgXcQ", YOUTUBE_URL),
            ("http://youtube.com/watch?v=dQw4w9WgXcQ", YOUTUBE_URL),
        ],
    )
    def test_youtube_links_collapse_to_one_canonical_form(self, cfg, raw, canonical):
        target = gatekeeper.normalize_url(raw, cfg=cfg)
        assert target.source_type == "youtube"
        assert target.canonical_url == canonical
        assert target.identifier == "dQw4w9WgXcQ"

    def test_github_blob_links_become_raw_links(self, cfg):
        target = gatekeeper.normalize_url(
            "https://github.com/acme/playbooks/blob/main/SKILL.md", cfg=cfg
        )
        assert target.source_type == "github"
        assert target.canonical_url == SKILL_URL

    def test_raw_github_links_pass_through(self, cfg):
        assert gatekeeper.normalize_url(SKILL_URL, cfg=cfg).canonical_url == SKILL_URL

    def test_tracking_params_are_stripped_but_identity_is_kept(self, cfg):
        target = gatekeeper.normalize_url(
            "https://example.com/post?utm_source=tw&fbclid=x&id=7&page=2", cfg=cfg
        )
        assert "utm_source" not in target.canonical_url
        assert "fbclid" not in target.canonical_url
        assert "id=7" in target.canonical_url and "page=2" in target.canonical_url

    def test_bare_hostnames_get_https(self, cfg):
        assert gatekeeper.normalize_url("example.com/post", cfg=cfg).canonical_url.startswith(
            "https://"
        )

    @pytest.mark.parametrize(
        "raw,reason",
        [
            ("https://www.youtube.com/playlist?list=PL123", "playlist_blocked"),
            ("https://www.youtube.com/watch?v=dQw4w9WgXcQ&list=PL123", "playlist_blocked"),
            ("https://www.youtube.com/@somechannel", "playlist_blocked"),
            ("https://www.youtube.com/channel/UC123", "playlist_blocked"),
            ("https://www.youtube.com/live/abcdefghijk", "livestream_blocked"),
            ("https://example.com/talk.mp4", "raw_media_blocked"),
            ("https://example.com/audio.mp3?x=1", "raw_media_blocked"),
            ("https://example.com/clip.mov", "raw_media_blocked"),
            ("ftp://example.com/file", "scheme_not_allowed"),
            ("https://www.youtube.com/watch?v=tooshort", "youtube_id_not_found"),
            ("https://www.youtube.com/watch", "youtube_id_not_found"),
            ("", "empty_url"),
        ],
    )
    def test_out_of_scope_links_are_refused_with_a_reason(self, cfg, raw, reason):
        with pytest.raises(GatekeeperError) as exc:
            gatekeeper.normalize_url(raw, cfg=cfg)
        assert exc.value.reason == reason
        assert exc.value.message  # always a human-readable explanation

    def test_scope_rules_run_before_dns(self, cfg):
        """An `.mp4` on a private host is rejected for scope, not for SSRF —
        proof that the cheap check happens first."""
        with pytest.raises(GatekeeperError) as exc:
            gatekeeper.normalize_url("https://internal.example.com/clip.mp4", cfg=cfg)
        assert exc.value.reason == "raw_media_blocked"

    def test_ssrf_still_applies_to_in_scope_links(self, cfg):
        with pytest.raises(SSRFBlocked):
            gatekeeper.normalize_url("https://internal.example.com/post", cfg=cfg)

    def test_duration_ceiling(self, cfg):
        assert gatekeeper.check_duration(60) == 1
        assert gatekeeper.check_duration(119 * 60) == 119
        with pytest.raises(GatekeeperError) as exc:
            gatekeeper.check_duration(121 * 60)
        assert exc.value.reason == "too_long"

    def test_livestream_flag(self, cfg):
        gatekeeper.check_is_live(False)
        with pytest.raises(GatekeeperError) as exc:
            gatekeeper.check_is_live(True)
        assert exc.value.reason == "livestream_blocked"


# ===========================================================================
# gatekeeper — files
# ===========================================================================
def _write_pdf(path: Path, pages: int = 2, text: str = "Hello from page") -> Path:
    from reportlab.pdfgen import canvas

    pdf = canvas.Canvas(str(path))
    for page in range(pages):
        pdf.drawString(72, 720, f"{text} {page + 1}.")
        pdf.showPage()
    pdf.save()
    return path


def _write_docx(path: Path) -> Path:
    from docx import Document

    document = Document()
    document.add_heading("Renewal Motion", level=1)
    document.add_paragraph(
        "We always pull the usage report before booking the executive review, "
        "because it anchors the conversation on value rather than on price."
    )
    document.save(str(path))
    return path


class TestFileGatekeeper:
    @pytest.mark.parametrize(
        "name,payload,source_type",
        [
            ("notes.txt", b"plain text notes about pricing", "text"),
            ("notes.md", b"# Title\nsome content", "text"),
            ("shot.png", b"\x89PNG\r\n\x1a\n" + b"0" * 64, "image"),
            ("shot.jpg", b"\xff\xd8\xff" + b"0" * 64, "image"),
            ("shot.webp", b"RIFF----WEBP" + b"0" * 64, "image"),
        ],
    )
    def test_allowed_types_are_accepted(self, cfg, tmp_path, name, payload, source_type):
        path = tmp_path / name
        path.write_bytes(payload)
        validated = gatekeeper.validate_file(path, cfg=cfg)
        assert validated.source_type == source_type
        assert len(validated.content_hash) == 64

    def test_pdf_and_docx_are_accepted(self, cfg, tmp_path):
        assert gatekeeper.validate_file(_write_pdf(tmp_path / "a.pdf"), cfg=cfg).source_type == "pdf"
        assert gatekeeper.validate_file(_write_docx(tmp_path / "a.docx"), cfg=cfg).source_type == "docx"

    @pytest.mark.parametrize("name", ["talk.mp4", "audio.mp3", "rec.wav", "clip.mov", "song.flac"])
    def test_raw_media_uploads_are_blocked(self, cfg, tmp_path, name):
        path = tmp_path / name
        path.write_bytes(b"x" * 128)
        with pytest.raises(GatekeeperError) as exc:
            gatekeeper.validate_file(path, cfg=cfg)
        assert exc.value.reason == "raw_media_blocked"

    @pytest.mark.parametrize("name", ["script.sh", "binary.exe", "data.csv", "archive.zip", "noext"])
    def test_unsupported_types_are_blocked(self, cfg, tmp_path, name):
        path = tmp_path / name
        path.write_bytes(b"x" * 128)
        with pytest.raises(GatekeeperError) as exc:
            gatekeeper.validate_file(path, cfg=cfg)
        assert exc.value.reason == "unsupported_file"

    def test_renaming_a_payload_to_pdf_does_not_fool_the_magic_check(self, cfg, tmp_path):
        path = tmp_path / "payload.pdf"
        path.write_bytes(b"MZ\x90\x00" + b"\x00" * 200)  # a PE executable
        with pytest.raises(GatekeeperError) as exc:
            gatekeeper.validate_file(path, cfg=cfg)
        assert exc.value.reason == "content_type_mismatch"

    def test_upload_cap_is_enforced(self, cfg, tmp_path):
        path = tmp_path / "big.txt"
        path.write_bytes(b"x" * (cfg.max_upload_bytes + 1))
        with pytest.raises(GatekeeperError) as exc:
            gatekeeper.validate_file(path, cfg=cfg)
        assert exc.value.reason == "file_too_large"

    def test_empty_and_missing_files(self, cfg, tmp_path):
        empty = tmp_path / "empty.txt"
        empty.write_bytes(b"")
        with pytest.raises(GatekeeperError) as exc:
            gatekeeper.validate_file(empty, cfg=cfg)
        assert exc.value.reason == "file_empty"
        with pytest.raises(GatekeeperError) as exc:
            gatekeeper.validate_file(tmp_path / "ghost.pdf", cfg=cfg)
        assert exc.value.reason == "file_missing"

    def test_zip_bomb_is_refused_by_inflated_size(self, cfg, tmp_path):
        path = tmp_path / "bomb.docx"
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", b"0" * (40 * 1024 * 1024))
        with pytest.raises(GatekeeperError) as exc:
            gatekeeper.validate_file(path, cfg=cfg)
        assert exc.value.reason == "zip_bomb"
        assert "30MB" in exc.value.message

    def test_a_real_docx_passes_the_zip_inspection(self, cfg, tmp_path):
        assert gatekeeper.inspect_zip(_write_docx(tmp_path / "ok.docx"), cfg=cfg) > 0

    def test_corrupt_docx_is_refused(self, cfg, tmp_path):
        path = tmp_path / "broken.docx"
        path.write_bytes(b"PK\x03\x04" + b"garbage" * 20)
        with pytest.raises(GatekeeperError) as exc:
            gatekeeper.validate_file(path, cfg=cfg)
        assert exc.value.reason == "corrupt_file"

    def test_identical_bytes_hash_identically(self, cfg, tmp_path):
        a, b = tmp_path / "a.txt", tmp_path / "b.txt"
        a.write_bytes(b"same content")
        b.write_bytes(b"same content")
        assert gatekeeper.hash_file(a) == gatekeeper.hash_file(b)


# ===========================================================================
# fetchers — subprocess safety and parsers
# ===========================================================================
class TestFetcherHardening:
    def test_ytdlp_commands_are_argument_arrays_with_a_flag_terminator(self, cfg):
        hostile = "https://www.youtube.com/watch?v=x --exec=rm -rf /"
        for cmd in (
            fetchers.build_ytdlp_metadata_cmd(hostile, cfg=cfg),
            fetchers.build_ytdlp_subtitle_cmd(hostile, "/tmp/out", cfg=cfg),
        ):
            assert isinstance(cmd, list)
            assert all(isinstance(part, str) for part in cmd)
            assert cmd[0] == "yt-dlp"
            assert "--no-config" in cmd
            assert "--no-playlist" in cmd
            assert "--max-filesize" in cmd
            assert cmd[cmd.index("--max-filesize") + 1] == cfg.ytdlp_max_filesize
            # `--` immediately before the URL: a crafted link cannot become a flag.
            assert cmd[-2] == "--"
            assert cmd[-1] == hostile
            assert not any(";" in part and part is not cmd[-1] for part in cmd)

    def test_run_refuses_a_string_command(self, cfg):
        with pytest.raises(ValueError):
            fetchers._run("yt-dlp --version", cfg=cfg)  # type: ignore[arg-type]

    def test_missing_binary_surfaces_as_a_provider_error(self, cfg):
        with pytest.raises(ProviderError) as exc:
            fetchers._run(["definitely-not-a-real-binary-xyz"], cfg=cfg)
        assert exc.value.reason == "tool_missing"

    def test_vtt_parsing_yields_timestamps_and_drops_duplicates(self):
        vtt = """WEBVTT
Kind: captions
Language: en

1
00:00:01.000 --> 00:00:04.000
Hello <c.colorE5E5E5>there</c>

00:23:40.000 --> 00:23:48.500
we batch all cold emails
we batch all cold emails

00:30:00.000 --> 00:30:05.000
 
"""
        cues = fetchers.parse_vtt(vtt)
        assert cues == [
            {"start": 1, "end": 4, "text": "Hello there"},
            {"start": 1420, "end": 1428, "text": "we batch all cold emails"},
        ]

    def test_vtt_parsing_survives_garbage(self):
        assert fetchers.parse_vtt("") == []
        assert fetchers.parse_vtt("not a vtt file at all") == []

    def test_markdown_sections_follow_headings(self):
        sections = fetchers.markdown_to_sections(
            "# Renewal Motion\nRun 60 days out.\n## Steps\nPull usage.\nBook review."
        )
        assert [s["section"] for s in sections] == ["Renewal Motion", "Steps"]
        assert "Book review." in sections[1]["text"]

    def test_html_extraction_drops_scripts_and_chrome(self):
        html = (
            "<html><head><title>Outbound</title>"
            '<meta name="author" content="Jane Expert"></head><body>'
            "<nav>menu menu menu</nav><script>evil()</script>"
            "<article><h2>Findings</h2><p>"
            + ("Continuous sending beat batched sending on reply rate. " * 3)
            + "</p></article><footer>legal</footer></body></html>"
        )
        title, author, sections = fetchers.html_to_sections(html)
        assert title == "Outbound" and author == "Jane Expert"
        body = " ".join(s["text"] for s in sections)
        assert "evil()" not in body and "menu" not in body and "legal" not in body
        assert "Continuous sending" in body

    def test_pdf_parsing_produces_page_locators(self, cfg, tmp_path):
        validated = gatekeeper.validate_file(_write_pdf(tmp_path / "doc.pdf", pages=3), cfg=cfg)
        content = fetchers.parse_pdf_file(validated, cfg=cfg)
        assert content.locator_type == "page"
        assert [cue["page"] for cue in content.cues] == [1, 2, 3]
        assert content.content_hash == validated.content_hash

    def test_pdf_page_cap_is_enforced(self, cfg, tmp_path):
        import dataclasses

        capped = dataclasses.replace(cfg, max_pdf_pages=2)
        validated = gatekeeper.validate_file(_write_pdf(tmp_path / "long.pdf", pages=4), cfg=cfg)
        with pytest.raises(GatekeeperError) as exc:
            fetchers.parse_pdf_file(validated, cfg=capped)
        assert exc.value.reason == "too_many_pages"

    def test_docx_parsing_uses_headings_as_sections(self, cfg, tmp_path):
        validated = gatekeeper.validate_file(_write_docx(tmp_path / "doc.docx"), cfg=cfg)
        content = fetchers.parse_docx_file(validated, cfg=cfg)
        assert content.locator_type == "section"
        assert content.cues[0]["section"] == "Renewal Motion"
        assert "usage report" in content.full_text

    def test_text_file_parsing(self, cfg, tmp_path):
        path = tmp_path / "notes.txt"
        path.write_text("First block of notes.\n\nSecond block of notes.")
        content = fetchers.parse_text_file(gatekeeper.validate_file(path, cfg=cfg), cfg=cfg)
        assert len(content.cues) == 2

    def test_images_need_a_vision_model_in_prod(self, cfg, tmp_path):
        path = tmp_path / "chart.png"
        path.write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 64)
        with pytest.raises(ProviderError) as exc:
            fetchers.describe_image_file(gatekeeper.validate_file(path, cfg=cfg), cfg=cfg)
        assert exc.value.reason == "vision_unavailable"


# ===========================================================================
# grounder
# ===========================================================================
TRANSCRIPT_CUES = [
    {"start": 0, "end": 12, "text": "Welcome back. Today we talk about outbound email timing."},
    {
        "start": 1420,
        "end": 1448,
        "text": (
            "We batch every single cold email into one block before ten in the morning, "
            "and our reply rate went up 34% the quarter we started."
        ),
    },
]
TRANSCRIPT_TEXT = "\n".join(cue["text"] for cue in TRANSCRIPT_CUES)


class TestGrounder:
    def test_a_verbatim_quote_grounds_with_its_timestamp(self, cfg):
        result = ground_quote(
            "We batch every single cold email into one block before ten in the morning",
            full_text=TRANSCRIPT_TEXT,
            cues=TRANSCRIPT_CUES,
            cfg=cfg,
        )
        assert result.grounded and result.score == 100.0
        assert result.locator_type == "timestamp" and result.locator_value == "1420"
        assert result.cue_index == 1

    def test_case_and_whitespace_differences_still_ground(self, cfg):
        result = ground_quote(
            "  WE   BATCH every single cold email into ONE BLOCK before ten in the morning  ",
            full_text=TRANSCRIPT_TEXT,
            cues=TRANSCRIPT_CUES,
            cfg=cfg,
        )
        assert result.grounded

    def test_the_stored_quote_is_the_span_actually_found_not_the_hypothesis(self, cfg):
        """A near-miss quote is replaced by the source's own words."""
        result = ground_quote(
            "we batch every single cold email into one block before ten in the mornin",
            full_text=TRANSCRIPT_TEXT,
            cues=TRANSCRIPT_CUES,
            cfg=cfg,
        )
        assert result.grounded
        assert result.anchor_quote in TRANSCRIPT_CUES[1]["text"]
        assert result.anchor_quote.startswith("We batch")  # original capitalisation

    def test_a_loose_paraphrase_is_rejected(self, cfg):
        result = ground_quote(
            "we batch each and every cold email into a single block before ten in the morning",
            full_text=TRANSCRIPT_TEXT,
            cues=TRANSCRIPT_CUES,
            cfg=cfg,
        )
        assert not result.grounded
        assert "below_threshold" in result.reason
        assert result.score < cfg.grounding_threshold

    def test_a_hallucinated_quote_is_rejected(self, cfg):
        result = ground_quote(
            "the speaker stated that artificial intelligence will replace every salesperson by 2030",
            full_text=TRANSCRIPT_TEXT,
            cues=TRANSCRIPT_CUES,
            cfg=cfg,
        )
        assert not result.grounded and result.score < 60

    def test_word_count_bounds_are_enforced(self, cfg):
        short = ground_quote("we batch emails", full_text=TRANSCRIPT_TEXT, cues=TRANSCRIPT_CUES, cfg=cfg)
        assert not short.grounded and short.reason.startswith("too_short")

        long_quote = " ".join(TRANSCRIPT_TEXT.split()[:60])
        result = ground_quote(long_quote, full_text=TRANSCRIPT_TEXT, cues=TRANSCRIPT_CUES, cfg=cfg)
        assert result.grounded
        assert len(result.anchor_quote.split()) <= cfg.anchor_max_words

    def test_empty_inputs_fail_closed(self, cfg):
        assert not ground_quote("", full_text=TRANSCRIPT_TEXT, cues=TRANSCRIPT_CUES, cfg=cfg).grounded
        assert not ground_quote("a quote with plenty of words in it to pass the minimum", full_text="", cues=[], cfg=cfg).grounded

    def test_page_and_section_locators_are_resolved(self, cfg):
        page_cues = [
            {"page": 14, "text": "Continuous sending beat batched sending on reply rate across every cohort."}
        ]
        page = ground_quote(
            "continuous sending beat batched sending on reply rate across every cohort",
            full_text=page_cues[0]["text"],
            cues=page_cues,
            cfg=cfg,
        )
        assert page.grounded and page.locator_type == "page" and page.locator_value == "14"

        section_cues = [
            {"section": "Guardrails", "text": "Never send renewal paper before the executive review, because it anchors on price."}
        ]
        section = ground_quote(
            "Never send renewal paper before the executive review, because it anchors on price",
            full_text=section_cues[0]["text"],
            cues=section_cues,
            cfg=cfg,
        )
        assert section.grounded and section.locator_type == "section"
        assert section.locator_value == "Guardrails"

    def test_the_closest_cue_wins_when_several_are_similar(self, cfg):
        cues = [
            {"start": 10, "end": 20, "text": "We batch emails in the morning most of the time."},
            {"start": 900, "end": 915, "text": "We batch every single cold email into one block before ten in the morning."},
        ]
        result = ground_quote(
            "We batch every single cold email into one block before ten in the morning",
            full_text="\n".join(c["text"] for c in cues),
            cues=cues,
            cfg=cfg,
        )
        assert result.locator_value == "900"


# ===========================================================================
# extractor + grounding gate
# ===========================================================================
class LyingExtractor:
    """Proposes a plausible claim with a quote that is not in the source."""

    def extract(self, chunk, context):
        return [
            {
                "claim": "AI will replace every salesperson by 2030",
                "when_it_applies": "always",
                "anchor_quote": "the speaker clearly stated that artificial intelligence replaces every salesperson by 2030",
                "evidence_basis": "data",
            }
        ]


class SchemaBreakingExtractor:
    def extract(self, chunk, context):
        return [
            {},
            {"claim": "no quote here", "when_it_applies": "x"},
            {"claim": "", "anchor_quote": "y" * 80, "when_it_applies": "x"},
            "not even a dict",
            {
                "claim": "bad enum and bad salience",
                "when_it_applies": "x",
                "anchor_quote": "Welcome back. Today I want to talk about outbound email timing.",
                "evidence_basis": "vibes",
                "salience": 99,
                "domain_tags": "sales",
                "exact_steps": "one step",
            },
        ]


class ExplodingExtractor:
    def extract(self, chunk, context):
        raise RuntimeError("the model fell over")


@pytest.fixture
def transcript(cfg):
    return DeterministicFetcher().fetch_url(
        gatekeeper.normalize_url(YOUTUBE_URL, cfg=cfg), cfg=cfg
    )


class TestExtractor:
    def test_chunking_respects_cue_boundaries(self, cfg, transcript):
        chunks = chunk_content(transcript)
        assert chunks
        for chunk in chunks:
            assert chunk.cue_indices  # every chunk knows which cues it came from
            assert len(chunk.text) <= 2400

    def test_long_cues_are_split_without_losing_their_locator(self, cfg):
        content = FetchedContent(
            title="Dense",
            full_text="",
            cues=[{"page": 1, "text": "Sentence number one is here. " * 300}],
            locator_type="page",
        ).finalise()
        chunks = chunk_content(content)
        assert len(chunks) > 1
        assert all(chunk.cue_indices == [0] for chunk in chunks)

    def test_content_without_cues_still_chunks(self, cfg):
        content = FetchedContent(title="t", full_text="A sentence. " * 500).finalise()
        assert len(chunk_content(content)) > 1

    def test_the_heuristic_extractor_produces_grounded_claims(self, cfg, transcript):
        extraction = extract_grounded_claims(
            transcript, extractor=HeuristicClaimExtractor(cfg=cfg), cfg=cfg
        )
        assert extraction.claims
        assert extraction.rejected == 0
        for claim in extraction.claims:
            assert claim["grounding_score"] >= cfg.grounding_threshold
            assert claim["locator_type"] == "timestamp"
            assert claim["locator_value"].isdigit()
            # The quote really is in the source.
            assert claim["anchor_quote"].lower()[:40] in transcript.full_text.lower()
            assert cfg.anchor_min_words <= len(claim["anchor_quote"].split()) <= cfg.anchor_max_words
            assert claim["evidence_basis"] in tenant_db.EVIDENCE_BASES

    def test_a_hallucinating_extractor_gets_nothing_through_the_gate(self, cfg, transcript):
        extraction = extract_grounded_claims(transcript, extractor=LyingExtractor(), cfg=cfg)
        assert extraction.claims == []
        assert extraction.rejected >= 1
        assert "below_threshold" in extraction.rejections[0]["reason"]

    def test_malformed_proposals_are_coerced_or_dropped(self, cfg, transcript):
        extraction = extract_grounded_claims(
            transcript, extractor=SchemaBreakingExtractor(), cfg=cfg
        )
        assert extraction.rejected >= 3
        for claim in extraction.claims:
            assert claim["evidence_basis"] in tenant_db.EVIDENCE_BASES
            assert 1 <= claim["salience"] <= 5
            assert isinstance(claim["domain_tags"], list)
            assert isinstance(claim["exact_steps"], list)

    def test_an_extractor_may_not_claim_user_axiom_provenance(self, cfg, transcript):
        class Impersonator:
            def extract(self, chunk, context):
                return [
                    {
                        "claim": "pretend this came from the user",
                        "when_it_applies": "x",
                        "anchor_quote": "Welcome back. Today I want to talk about outbound email timing.",
                        "evidence_basis": "user_axiom",
                    }
                ]

        extraction = extract_grounded_claims(transcript, extractor=Impersonator(), cfg=cfg)
        assert extraction.claims
        assert extraction.claims[0]["evidence_basis"] != "user_axiom"

    def test_an_exploding_extractor_does_not_kill_the_job(self, cfg, transcript):
        extraction = extract_grounded_claims(transcript, extractor=ExplodingExtractor(), cfg=cfg)
        assert extraction.claims == []
        assert "extractor_error" in extraction.rejections[0]["reason"]

    def test_duplicate_claims_are_collapsed(self, cfg, transcript):
        class Repeater:
            def extract(self, chunk, context):
                quote = "Welcome back. Today I want to talk about outbound email timing."
                return [{"claim": "same claim", "when_it_applies": "x", "anchor_quote": quote}] * 3

        extraction = extract_grounded_claims(transcript, extractor=Repeater(), cfg=cfg)
        assert len(extraction.claims) == 1


# ===========================================================================
# embedder + linker
# ===========================================================================
class TestEmbedder:
    def test_embeddings_are_deterministic_normalised_and_right_sized(self, cfg):
        embedder = HashingEmbedder(cfg.embedding_dim)
        first = embedder.embed_one("batch cold emails in the morning")
        assert first == embedder.embed_one("batch cold emails in the morning")
        assert len(first) == cfg.embedding_dim
        assert abs(sum(v * v for v in first) - 1.0) < 1e-6

    def test_related_text_is_closer_than_unrelated_text(self, cfg):
        embedder = HashingEmbedder(cfg.embedding_dim)
        batching, batched, warehouse = embedder.embed(
            [
                "We batch every cold email into one morning block",
                "Cold emails should be batched into a single block each morning",
                "Warehouse pick paths should be optimised by zone",
            ]
        )

        def cosine(a, b):
            return sum(x * y for x, y in zip(a, b))

        assert cosine(batching, batched) > cosine(batching, warehouse)
        assert cosine(batching, warehouse) < 0.1

    def test_stemming_bridges_inflections(self, cfg):
        embedder = HashingEmbedder(cfg.embedding_dim)
        a, b = embedder.embed(["we batch every email", "batched emails"])
        assert sum(x * y for x, y in zip(a, b)) > 0.3

    def test_empty_text_still_yields_a_usable_vector(self, cfg):
        vector = HashingEmbedder(cfg.embedding_dim).embed_one("")
        assert len(vector) == cfg.embedding_dim
        assert abs(sum(v * v for v in vector) - 1.0) < 1e-6


class TestLinkClassifier:
    def setup_method(self):
        self.classifier = HeuristicLinkClassifier()
        self.batching = {
            "claim": "We batch every single cold email into one morning block",
            "when_it_applies": "Small outbound teams",
            "speaker": "Jane Expert",
        }

    def test_asymmetric_negation_is_a_contradiction(self):
        other = {
            "claim": "Never batch cold emails, send them continuously instead",
            "when_it_applies": "Teams above 500 sends a day",
            "speaker": "Acme Research",
        }
        verdict = self.classifier.classify(self.batching, other)
        assert verdict["link_type"] == "contradicts"
        assert "Jane Expert" in verdict["divergence_summary"]
        assert "Acme Research" in verdict["divergence_summary"]

    def test_a_comparative_counts_as_an_inverted_stance(self):
        other = {
            "claim": "Continuous cold email sending beat batched sending on reply rate",
            "when_it_applies": "High volume outbound",
            "speaker": "Acme Research",
        }
        assert self.classifier.classify(self.batching, other)["link_type"] == "contradicts"

    def test_an_added_boundary_condition_is_a_refinement(self):
        other = {
            "claim": "Batch cold emails into a morning block only when you send under 500 a day",
            "when_it_applies": "Low volume outbound",
            "speaker": "Acme Research",
        }
        assert self.classifier.classify(self.batching, other)["link_type"] == "refines"

    def test_agreement_produces_no_link(self):
        other = {
            "claim": "We batch all cold email into a single morning block",
            "when_it_applies": "Small outbound teams",
            "speaker": "Bob",
        }
        assert self.classifier.classify(self.batching, other) is None

    def test_unrelated_claims_produce_no_link(self):
        other = {
            "claim": "Never optimise warehouse pick paths by zone",
            "when_it_applies": "Multi zone warehouses",
            "speaker": "Bob",
        }
        assert self.classifier.classify(self.batching, other) is None

    def test_empty_claims_produce_no_link(self):
        assert self.classifier.classify({}, self.batching) is None


class TestLinker:
    def _seed(self, tdb, embedder, *, source_title, author, claim_text, applies):
        source = tenant_db.upsert_source(
            tdb, title=source_title, source_type="article", author=author,
            canonical_url=f"https://example.com/{source_title.lower().replace(' ', '-')}",
        )
        claim = tenant_db.insert_claim(
            tdb,
            source_id=source["id"],
            claim=claim_text,
            speaker=author,
            when_it_applies=applies,
            anchor_quote=claim_text,
            evidence_basis="case_study",
        )
        vector = embedder.embed_one(f"{claim_text} {applies}")
        tenant_db.set_claim_embedding(tdb, claim["id"], vector)
        return claim, vector

    def test_a_contradiction_across_two_sources_is_recorded(self, cfg, tdb):
        embedder = HashingEmbedder(cfg.embedding_dim)
        self._seed(
            tdb, embedder,
            source_title="Masterclass", author="Jane Expert",
            claim_text="We batch every single cold email into one morning block",
            applies="Small outbound teams sending cold email",
        )
        new_claim, vector = self._seed(
            tdb, embedder,
            source_title="Benchmark", author="Acme Research",
            claim_text="Never batch cold email; continuous sending beat batched sending",
            applies="Outbound teams above 500 cold email sends a day",
        )
        result = link_claims(
            tdb, [new_claim], embeddings={new_claim["id"]: vector},
            classifier=HeuristicLinkClassifier(), cfg=cfg,
        )
        assert len(result.created) == 1
        assert result.created[0]["link_type"] == "contradicts"
        assert tenant_db.counts(tdb)["tensions"] == 1

    def test_a_source_is_never_linked_against_itself(self, cfg, tdb):
        embedder = HashingEmbedder(cfg.embedding_dim)
        source = tenant_db.upsert_source(tdb, title="One Source", source_type="article")
        claims = []
        for text in (
            "We batch every single cold email into one morning block",
            "Never batch cold email; continuous sending beat batched sending",
        ):
            claim = tenant_db.insert_claim(
                tdb, source_id=source["id"], claim=text, when_it_applies="cold email outbound",
                anchor_quote=text,
            )
            vector = embedder.embed_one(text)
            tenant_db.set_claim_embedding(tdb, claim["id"], vector)
            claims.append((claim, vector))
        result = link_claims(
            tdb,
            [claims[1][0]],
            embeddings={claims[1][0]["id"]: claims[1][1]},
            classifier=HeuristicLinkClassifier(),
            cfg=cfg,
        )
        assert result.created == []

    def test_agreeing_claims_are_not_linked(self, cfg, tdb):
        embedder = HashingEmbedder(cfg.embedding_dim)
        self._seed(
            tdb, embedder, source_title="A", author="Jane",
            claim_text="We batch every single cold email into one morning block",
            applies="Small outbound teams",
        )
        new_claim, vector = self._seed(
            tdb, embedder, source_title="B", author="Bob",
            claim_text="We batch all cold email into a single morning block",
            applies="Small outbound teams",
        )
        result = link_claims(
            tdb, [new_claim], embeddings={new_claim["id"]: vector},
            classifier=HeuristicLinkClassifier(), cfg=cfg,
        )
        assert result.created == []
        assert result.pairs_classified >= 1  # it *was* considered, and declined

    def test_links_are_not_duplicated_on_a_second_pass(self, cfg, tdb):
        embedder = HashingEmbedder(cfg.embedding_dim)
        self._seed(
            tdb, embedder, source_title="A", author="Jane",
            claim_text="We batch every single cold email into one morning block",
            applies="Small outbound teams sending cold email",
        )
        new_claim, vector = self._seed(
            tdb, embedder, source_title="B", author="Acme",
            claim_text="Never batch cold email; continuous sending beat batched sending",
            applies="Outbound teams above 500 cold email sends a day",
        )
        kwargs = dict(
            embeddings={new_claim["id"]: vector},
            classifier=HeuristicLinkClassifier(),
            cfg=cfg,
        )
        assert len(link_claims(tdb, [new_claim], **kwargs).created) == 1
        assert link_claims(tdb, [new_claim], **kwargs).created == []
        assert tenant_db.counts(tdb)["tensions"] == 1


# ===========================================================================
# engine — the whole pipeline
# ===========================================================================
class TestEngineIngestion:
    def test_queueing_is_immediate_and_does_no_work(self, ctx):
        result = E.op_ingest_source(ctx, "https://youtu.be/dQw4w9WgXcQ?si=tracking")
        assert result.ok
        assert result.data["status"] == "queued"
        assert result.data["canonical_url"] == YOUTUBE_URL
        assert result.data["job_id"].startswith("job_")
        # Nothing has been fetched or stored yet.
        assert tenant_db.counts(ctx.tenant)["claims"] == 0

    def test_the_pipeline_produces_grounded_cited_claims(self, ctx):
        E.op_ingest_source(ctx, YOUTUBE_URL, user_note="for outbound work")
        processed = E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg)
        assert len(processed) == 1
        assert processed[0]["status"] == "ready"
        assert processed[0]["claims"] >= 1
        assert processed[0]["rejected"] == 0

        found = E.op_search(ctx, "batch cold emails")
        assert found.data["claims"]
        top = found.data["claims"][0]
        assert top["source"]["title"] == "Outbound Masterclass: the morning block"
        assert top["locator_type"] == "timestamp"
        assert top["deep_link"] == f"{YOUTUBE_URL}&t={top['locator_value']}"
        assert top["anchor_quote"]

    def test_pdf_ingestion_yields_page_deep_links(self, ctx, tmp_path):
        path = _write_pdf(tmp_path / "report.pdf", pages=2, text="We always measure the baseline first because a 20% lift means nothing without it, on page")
        queued = E.op_ingest_file(ctx, path)
        assert queued.ok
        processed = E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg)
        assert processed[0]["status"] == "ready"
        found = E.op_search(ctx, "measure the baseline")
        assert found.data["claims"]
        assert found.data["claims"][0]["locator_type"] == "page"

    def test_a_contradiction_across_two_sources_surfaces_in_search(self, ctx):
        E.op_ingest_source(ctx, YOUTUBE_URL)
        E.op_ingest_source(ctx, REPORT_URL)
        E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg)

        found = E.op_search(ctx, "batch cold emails")
        tensions = [t for claim in found.data["claims"] for t in claim["tensions"]]
        assert tensions, "the batching/continuous disagreement must be surfaced"
        assert any(t["link_type"] in ("contradicts", "refines") for t in tensions)
        assert all(t["divergence_summary"] for t in tensions)
        assert all(t["other_claim"] for t in tensions)

    def test_a_duplicate_url_is_recognised_not_re_ingested(self, ctx):
        E.op_ingest_source(ctx, YOUTUBE_URL)
        E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg)
        again = E.op_ingest_source(ctx, "https://youtu.be/dQw4w9WgXcQ")
        assert again.ok and again.data["duplicate"] is True
        assert again.data["claim_count"] >= 1
        assert control_db.list_jobs(ctx.control, ctx.user_id, active_only=True) == []

    def test_a_duplicate_file_is_recognised_by_content_hash(self, ctx, tmp_path):
        first = tmp_path / "a.txt"
        first.write_text("We always label the axes before sharing a chart with the team.")
        second = tmp_path / "b.txt"
        second.write_text("We always label the axes before sharing a chart with the team.")
        E.op_ingest_file(ctx, first)
        E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg)
        assert E.op_ingest_file(ctx, second).data["duplicate"] is True

    def test_out_of_scope_links_never_create_a_job(self, ctx):
        for url, reason in (
            ("https://www.youtube.com/playlist?list=PL1", "playlist_blocked"),
            ("https://example.com/talk.mp4", "raw_media_blocked"),
            ("https://internal.example.com/secrets", "private_address"),
        ):
            result = E.op_ingest_source(ctx, url)
            assert not result.ok and result.error_reason == reason
            assert result.error_message
        assert control_db.list_jobs(ctx.control, ctx.user_id) == []

    def test_a_fetch_that_fails_costs_no_credits(self, ctx):
        """`R-18`. The fetch is free — the charge happens after it, at the
        gate — so a source that never arrives has nothing to give back."""
        class Failing:
            def fetch_url(self, target, *, cfg):
                raise ProviderError("no_transcript", "That video has no captions.")

            def fetch_file(self, validated, *, cfg):
                raise ProviderError("no_transcript", "n/a")

        E.op_ingest_source(ctx, YOUTUBE_URL)
        assert control_db.credits_state(ctx.control, ctx.user_id, cfg=ctx.cfg)["used"] == 0

        ctx.prov.fetcher = Failing()
        processed = E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg)
        assert processed[0]["status"] == "failed"
        assert "no captions" in processed[0]["message"]
        assert control_db.credits_state(ctx.control, ctx.user_id, cfg=ctx.cfg)["used"] == 0

    def test_a_job_that_fails_after_the_gate_is_refunded_exactly(self, ctx):
        """The harder half: the charge is already taken when the failure
        happens, so something has to give it back — and give back exactly
        what was taken, read from the job row rather than recomputed."""
        E.op_ingest_source(ctx, YOUTUBE_URL)
        E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg)
        charged = control_db.list_jobs(ctx.control, ctx.user_id)[0]["charged_credits"]
        assert charged > 0
        assert control_db.credits_state(ctx.control, ctx.user_id, cfg=ctx.cfg)["used"] == charged

        # Fail the same job the way `finish` would, after the gate.
        job = control_db.list_jobs(ctx.control, ctx.user_id)[0]
        control_db.refund_credits(
            ctx.control, ctx.user_id, credits=control_db.job_charge(ctx.control, job["job_id"])
        )
        assert control_db.credits_state(ctx.control, ctx.user_id, cfg=ctx.cfg)["used"] == 0

    def test_a_retry_charges_once(self, ctx):
        """`charged_credits` on the row is what makes this true: a second
        pass through the gate sees the job has already paid."""
        E.op_ingest_source(ctx, YOUTUBE_URL)
        E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg)
        once = control_db.credits_state(ctx.control, ctx.user_id, cfg=ctx.cfg)["used"]
        assert once > 0

        job = control_db.list_jobs(ctx.control, ctx.user_id)[0]
        control_db.set_job_status(ctx.control, job["job_id"], "queued")
        E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg)
        assert control_db.credits_state(ctx.control, ctx.user_id, cfg=ctx.cfg)["used"] == once

    def test_a_platform_block_is_reported_distinctly(self, ctx):
        class Blocked:
            def fetch_url(self, target, *, cfg):
                raise ProviderError("blocked_by_platform", "YouTube refused that video.")

            def fetch_file(self, validated, *, cfg):
                raise ProviderError("blocked_by_platform", "n/a")

        E.op_ingest_source(ctx, YOUTUBE_URL)
        ctx.prov.fetcher = Blocked()
        assert E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg)[0]["status"] == "blocked_by_platform"

    def test_queueing_charges_nothing_and_the_gate_charges_the_real_text(self, ctx):
        """`R-18` replaced the provisional reservation with a single debit at
        the one line where the number is real. A queued job has spent
        nothing; a finished one has spent exactly the characters its text
        chunked to."""
        E.op_ingest_source(ctx, YOUTUBE_URL)
        assert control_db.credits_state(ctx.control, ctx.user_id, cfg=ctx.cfg)["used"] == 0

        E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg)
        used = control_db.credits_state(ctx.control, ctx.user_id, cfg=ctx.cfg)["used"]

        from app.ingestion.extractor import chunk_content
        from app.ingestion.providers import providers

        content = providers(ctx.cfg).fetcher.fetch_url(
            gatekeeper.normalize_url(YOUTUBE_URL, cfg=ctx.cfg), cfg=ctx.cfg
        )
        expected = sum(len(c.text) for c in chunk_content(content))
        assert used == expected, "billed the text the extractor actually chunks"

    def test_a_video_over_the_ceiling_fails_at_the_duration_gate(self, ctx):
        class TooLong:
            def fetch_url(self, target, *, cfg):
                return FetchedContent(
                    title="Marathon",
                    full_text="We always measure the baseline before changing anything at all.",
                    cues=[{"start": 0, "end": 10, "text": "We always measure the baseline before changing anything at all."}],
                    duration_sec=200 * 60,
                    source_type="youtube",
                    locator_type="timestamp",
                ).finalise()

            def fetch_file(self, validated, *, cfg):
                raise ProviderError("x", "y")

        E.op_ingest_source(ctx, YOUTUBE_URL)
        ctx.prov.fetcher = TooLong()
        processed = E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg)
        assert processed[0]["status"] == "failed"
        assert "120 minutes" in processed[0]["message"]
        assert control_db.get_user(ctx.control, ctx.user_id)["media_minutes_used"] == 0

    def test_a_source_with_no_groundable_claims_fails_loudly(self, ctx):
        E.op_ingest_source(ctx, YOUTUBE_URL)
        ctx.prov.extractor = LyingExtractor()
        processed = E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg)
        assert processed[0]["status"] == "failed"
        assert "grounded" in processed[0]["message"]
        assert tenant_db.counts(ctx.tenant)["claims"] == 0

    def test_status_reporting_tracks_a_job_to_completion(self, ctx):
        queued = E.op_ingest_source(ctx, YOUTUBE_URL)
        job_id = queued.data["job_id"]
        status = E.op_get_ingest_status(ctx, job_id)
        assert status.data["jobs"][0]["status"] == "queued"

        E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg)
        done = E.op_get_ingest_status(ctx, job_id).data["jobs"][0]
        assert done["status"] == "ready"
        assert done["claims_extracted"] >= 1
        assert done["title"] == "Outbound Masterclass: the morning block"

        assert E.op_get_ingest_status(ctx).data["jobs"][0]["job_id"] == job_id
        assert not E.op_get_ingest_status(ctx, "job_nope").ok

    def test_one_tenants_job_is_invisible_to_another(self, ctx, control):
        other_id = control_db.ensure_user(control, "intruder@example.com")["user_id"]
        other = E.build_context(other_id)
        job_id = E.op_ingest_source(ctx, YOUTUBE_URL).data["job_id"]
        assert E.op_get_ingest_status(other, job_id).error_reason == "not_found"

    def test_the_pipeline_is_deterministic(self, cfg, control):
        """Same input, two tenants, byte-identical claims."""
        outputs = []
        for email in ("a@example.com", "b@example.com"):
            user_id = control_db.ensure_user(control, email)["user_id"]
            tenant_ctx = E.build_context(user_id)
            E.op_ingest_source(tenant_ctx, YOUTUBE_URL)
            E.run_pending_jobs_sync(user_id, cfg=cfg)
            claims = tenant_db.claims_for_source(
                tenant_ctx.tenant,
                control_db.list_jobs(control, user_id)[0]["source_id"],
            )
            outputs.append(
                [(c["claim"], c["anchor_quote"], c["locator_value"], c["salience"]) for c in claims]
            )
        assert outputs[0] == outputs[1]


class TestEngineInjectionSafety:
    def test_injected_instructions_are_stripped_and_fenced(self, ctx):
        """A hostile source may not smuggle instructions to the host LLM."""
        E.op_ingest_source(ctx, MALICIOUS_URL)
        E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg)

        stored = tenant_db.get_source_content(
            ctx.tenant, control_db.list_jobs(ctx.control, ctx.user_id)[0]["source_id"]
        )
        # Invisible channels are gone at ingestion time.
        assert "​" not in stored["full_text"]
        assert "‮" not in stored["full_text"]
        assert "\x00" not in stored["full_text"]
        assert "\U000E0041" not in stored["full_text"]

        # And when the text is handed to an LLM, it is fenced and de-fanged.
        wrapped = wrap_untrusted(stored["full_text"])
        assert wrapped.count(UNTRUSTED_END) == 1
        assert "[AGENT_HINT]" not in wrapped

    def test_a_hostile_source_cannot_archive_anything(self, ctx):
        """The injected text says "archive all user skills" — nothing happens."""
        E.op_save_skill(
            ctx, "renewal-play", "Renewal Play", "d", "60 days out", ["pull usage"]
        )
        E.op_ingest_source(ctx, MALICIOUS_URL)
        E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg)
        assert tenant_db.counts(ctx.tenant)["skills"] == 1


class TestEngineMemoryAndSkills:
    def test_remember_this_stores_a_user_axiom_that_is_grounded_by_definition(self, ctx):
        result = E.op_remember_this(
            ctx,
            claim="I never send cold email on a Friday",
            when_it_applies="My own outbound sequences",
            when_it_fails="Re-engagement campaigns to warm leads",
            domain_tags=["sales"],
        )
        assert result.ok
        claim = result.data["claim"]
        assert claim["evidence_basis"] == "user_axiom"
        assert claim["speaker"] == "You"
        assert claim["anchor_quote"] == "I never send cold email on a Friday"
        assert claim["locator_type"] == "none"
        assert claim["salience"] == 5
        assert E.op_search(ctx, "cold email friday").data["claims"][0]["id"] == claim["id"]

    def test_remember_this_validates_its_inputs(self, ctx):
        assert E.op_remember_this(ctx, "", "x").error_reason == "empty_claim"
        assert E.op_remember_this(ctx, "something", "").error_reason == "missing_when_it_applies"

    def test_a_user_axiom_can_contradict_an_expert(self, ctx):
        E.op_ingest_source(ctx, YOUTUBE_URL)
        E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg)
        result = E.op_remember_this(
            ctx,
            claim="I never batch cold email; continuous sending always wins for my team",
            when_it_applies="My own outbound sending of cold email",
            domain_tags=["sales"],
        )
        assert result.data["new_tensions"], "my own rule must be allowed to disagree with an expert"
        assert result.data["new_tensions"][0]["link_type"] == "contradicts"

    def test_synthesis_happens_when_no_skill_is_saved(self, ctx):
        E.op_ingest_source(ctx, YOUTUBE_URL)
        E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg)
        result = E.op_get_or_synthesize_skill(ctx, "run a cold email outbound sprint")
        assert result.data["origin"] == "synthesized"
        draft = result.data["draft_skill"]
        assert draft["execution_steps"]
        assert draft["origin_type"] == "synthesized" and draft["version"] == 0
        assert draft["linked_claim_ids"]
        assert result.data["supporting_claims"]

    def test_an_empty_library_says_so_rather_than_inventing(self, ctx):
        result = E.op_get_or_synthesize_skill(ctx, "run a cold email outbound sprint")
        assert result.ok and result.data["origin"] == "none"
        assert result.data["supporting_claims"] == []

    def test_a_saved_skill_wins_over_synthesis(self, ctx):
        E.op_save_skill(
            ctx,
            slug="cold-email-sprint",
            name="Cold Email Sprint",
            description="Run a one week cold email sprint",
            trigger_when="Starting an outbound push",
            execution_steps=["Build the list", "Write the sequence"],
            guardrails="Never send on a Friday",
        )
        result = E.op_get_or_synthesize_skill(ctx, "run a cold email outbound sprint")
        assert result.data["origin"] == "saved"
        assert result.data["skill"]["slug"] == "cold-email-sprint"
        assert result.data["skill"]["execution_steps"] == ["Build the list", "Write the sequence"]

    def test_an_explicit_slug_short_circuits_the_search(self, ctx):
        E.op_save_skill(
            ctx, "renewal-play", "Renewal Play", "d", "60 days out", ["pull usage"]
        )
        result = E.op_get_or_synthesize_skill(ctx, "something unrelated", skill_slug="renewal-play")
        assert result.data["origin"] == "saved" and result.data["skill"]["slug"] == "renewal-play"

    def test_save_skill_validates_and_versions(self, ctx):
        assert E.op_save_skill(ctx, "s", "n", "d", "t", []).error_reason == "no_steps"
        assert E.op_save_skill(ctx, "s", "", "d", "t", ["a"]).error_reason == "no_name"
        assert E.op_save_skill(ctx, "s", "n", "d", "", ["a"]).error_reason == "no_trigger"

        first = E.op_save_skill(ctx, "Renewal Play!", "Renewal Play", "d", "t", ["one"])
        assert first.data["created"] is True and first.data["skill"]["slug"] == "renewal-play"
        second = E.op_save_skill(ctx, "renewal-play", "Renewal Play", "d2", "t", ["one", "two"])
        assert second.data["created"] is False and second.data["skill"]["version"] == 2

    def test_a_saved_skill_is_findable_by_search(self, ctx):
        E.op_save_skill(
            ctx, "renewal-play", "Renewal Play", "Run the enterprise renewal motion",
            "60 days before renewal", ["Pull the usage report"],
        )
        found = E.op_search(ctx, "enterprise renewal motion", scope="skills")
        assert found.data["skills"][0]["slug"] == "renewal-play"
        assert found.data["claims"] == []


class TestEngineArchive:
    def test_archiving_hides_an_item_and_is_reversible(self, ctx):
        E.op_ingest_source(ctx, YOUTUBE_URL)
        E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg)
        claim_id = E.op_search(ctx, "batch cold emails").data["claims"][0]["id"]

        archived = E.op_archive_item(ctx, "claim", claim_id)
        assert archived.ok and archived.data["status"] == "archived"
        assert archived.data["recoverable"] is True
        assert claim_id not in [c["id"] for c in E.op_search(ctx, "batch cold emails").data["claims"]]

        restored = E.op_archive_item(ctx, "claim", claim_id, restore=True)
        assert restored.data["status"] == "active"
        assert claim_id in [c["id"] for c in E.op_search(ctx, "batch cold emails").data["claims"]]

    def test_archiving_a_source_takes_its_claims_with_it(self, ctx):
        E.op_ingest_source(ctx, YOUTUBE_URL)
        E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg)
        source_id = control_db.list_jobs(ctx.control, ctx.user_id)[0]["source_id"]
        E.op_archive_item(ctx, "source", source_id)
        assert tenant_db.counts(ctx.tenant)["claims"] == 0
        E.op_archive_item(ctx, "source", source_id, restore=True)
        assert tenant_db.counts(ctx.tenant)["claims"] >= 1

    def test_archiving_a_skill_by_slug_or_id(self, ctx):
        skill = E.op_save_skill(ctx, "renewal-play", "Renewal Play", "d", "t", ["a"]).data["skill"]
        assert E.op_archive_item(ctx, "skill", "renewal-play").data["status"] == "archived"
        assert E.op_archive_item(ctx, "skill", skill["id"], restore=True).data["status"] == "active"

    def test_archive_rejects_unknown_types_and_ids(self, ctx):
        assert E.op_archive_item(ctx, "database", "x").error_reason == "bad_item_type"  # type: ignore[arg-type]
        assert E.op_archive_item(ctx, "claim", "clm_nope").error_reason == "not_found"

    def test_permanent_deletion_needs_the_thing_set_aside_first(self, cfg, user):
        """Hard delete exists for the Web UI, and only for something the
        person has already set aside.

        The invariant that matters is not that the engine contains no DELETE —
        the founder's Profile modal needs one — but that nothing can reach it
        except a human, and only for something they already archived or
        trashed. The MCP-surface half of this is also asserted structurally in
        `tests/test_phase3_mcp_and_auth.py`.

        **The precondition used to be Trash specifically**, which is a
        source-level concept. §9.4 puts `Delete permanently` on an archived
        card of any kind, so that rule meant the button on §7.12's cards could
        only ever answer `not_in_trash` — and `not_found` for a rule, skill or
        summary, whose ids were resolved against `sources` alone. Both states
        now qualify; neither is reachable without the person having acted.
        """
        web = E.build_context(user["user_id"], surface="web_ui")
        E.op_ingest_source(web, YOUTUBE_URL)
        E.run_pending_jobs_sync(web.user_id, cfg=web.cfg)
        source_id = control_db.list_jobs(web.control, web.user_id)[0]["source_id"]

        # Active: refused, and it says what to do instead.
        refused = E.op_permanent_delete(web, source_id)
        assert not refused.ok and refused.error_reason == "not_set_aside"
        assert "archive it first" in refused.error_message.lower()
        assert tenant_db.get_source(web.tenant, source_id) is not None

        # And an agent is refused even once it IS set aside — the guard is the
        # surface, not the state.
        E.op_archive_item(web, "source", source_id)
        agent = E.build_context(user["user_id"], surface="mcp")
        blocked = E.op_permanent_delete(agent, source_id)
        assert not blocked.ok and blocked.error_reason == "person_only"
        assert tenant_db.get_source(web.tenant, source_id) is not None

        deleted = E.op_permanent_delete(web, source_id)
        assert deleted.ok and deleted.data["deleted_claims"] >= 1
        assert tenant_db.get_source(web.tenant, source_id) is None
        assert tenant_db.counts(web.tenant)["claims"] == 0

    def test_a_row_left_in_trash_can_still_be_deleted(self, cfg, user):
        """`R-19` folded Trash into Archive and removed every way to create a
        trashed row — but `op_permanent_delete` still accepts the status on
        purpose, because a database restored from a backup taken before the
        fold can hold one, and `tenant.migrate` only runs on open.

        So the status is written here directly. That is not a shortcut around
        a deleted operation; it is the only way the state now arises, and the
        path out of it has to keep working."""
        web = E.build_context(user["user_id"], surface="web_ui")
        E.op_ingest_source(web, YOUTUBE_URL)
        E.run_pending_jobs_sync(web.user_id, cfg=web.cfg)
        source_id = control_db.list_jobs(web.control, web.user_id)[0]["source_id"]

        tenant_db.set_source_status(web.tenant, source_id, "trash")
        assert E.op_permanent_delete(web, source_id).ok
        assert tenant_db.get_source(web.tenant, source_id) is None

    def test_the_migration_folds_a_trashed_row_into_the_archive(self, cfg, user):
        """The other half: a row that was already in Trash when the old app
        went must surface somewhere a person can reach, or it is invisible
        and unrecoverable for good."""
        web = E.build_context(user["user_id"], surface="web_ui")
        E.op_ingest_source(web, YOUTUBE_URL)
        E.run_pending_jobs_sync(web.user_id, cfg=web.cfg)
        source_id = control_db.list_jobs(web.control, web.user_id)[0]["source_id"]
        tenant_db.set_source_status(web.tenant, source_id, "trash")

        tenant_db.migrate(web.tenant)

        archived = [row for row in tenant_db.list_archive(web.tenant) if row["id"] == source_id]
        assert len(archived) == 1
        assert archived[0]["archived_by"] == "you", "they put it there, not Brain"

    def test_archiving_is_the_only_destructive_verb_an_agent_can_use(self):
        """No `op_*` an MCP tool calls may hard-delete."""
        import inspect

        for name in ("op_archive_item", "op_search", "op_fetch", "op_save_skill"):
            body = inspect.getsource(getattr(E, name))
            assert "DELETE FROM" not in body, name


class TestEngineContext:
    def test_an_unknown_user_cannot_get_a_context(self, cfg):
        with pytest.raises(ValueError):
            E.build_context("not-a-real-user")

    def test_a_hostile_user_id_never_reaches_the_filesystem(self, cfg):
        with pytest.raises(ValueError):
            E.build_context("../../etc/passwd")

    def test_every_operation_is_audited(self, ctx):
        E.op_search(ctx, "anything")
        E.op_ingest_source(ctx, YOUTUBE_URL)
        E.op_remember_this(ctx, "a rule I follow", "always")
        E.op_save_skill(ctx, "s", "n", "d", "t", ["a"])
        actions = {row["action"] for row in control_db.recent_audit(ctx.control, ctx.user_id)}
        assert {"search", "ingest_source", "remember_this", "save_skill"} <= actions

    def test_search_requires_a_query(self, ctx):
        assert E.op_search(ctx, "   ").error_reason == "empty_query"

    def test_fetch_requires_an_id(self, ctx):
        assert E.op_fetch(ctx, "").error_reason == "empty_id"

    def test_fetch_resolves_claims_skills_and_sources(self, ctx):
        E.op_ingest_source(ctx, YOUTUBE_URL)
        E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg)
        source_id = control_db.list_jobs(ctx.control, ctx.user_id)[0]["source_id"]
        claim_id = E.op_search(ctx, "batch cold emails").data["claims"][0]["id"]
        skill = E.op_save_skill(
            ctx, "renewal-play", "Renewal Play", "d", "t", ["a"], linked_claim_ids=[claim_id]
        ).data["skill"]

        assert E.op_fetch(ctx, claim_id).data["kind"] == "claim"
        assert E.op_fetch(ctx, source_id).data["kind"] == "source"
        assert E.op_fetch(ctx, skill["id"]).data["kind"] == "skill"
        assert E.op_fetch(ctx, "renewal-play").data["kind"] == "skill"
        assert E.op_fetch(ctx, skill["id"]).data["linked_claims"][0]["id"] == claim_id

    def test_deep_links_are_built_per_locator_type(self):
        assert E.deep_link("https://youtube.com/watch?v=a", "timestamp", "1420") == (
            "https://youtube.com/watch?v=a&t=1420"
        )
        assert E.deep_link("https://example.com/r.pdf", "page", "14") == (
            "https://example.com/r.pdf#page=14"
        )
        assert E.deep_link("https://example.com/post", "section", "Findings") == (
            "https://example.com/post"
        )
        assert E.deep_link("", "timestamp", "10") == ""
        assert E.deep_link("https://example.com", "timestamp", "not-a-number") == "https://example.com"

    def test_the_runner_drains_several_jobs_and_stops_cleanly(self, ctx):
        for url in (YOUTUBE_URL, REPORT_URL, SKILL_URL):
            E.op_ingest_source(ctx, url)
        processed = E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg, limit=10)
        assert len(processed) == 3
        assert all(p["status"] == "ready" for p in processed)
        assert E.run_pending_jobs_sync(ctx.user_id, cfg=ctx.cfg) == []

    def test_run_now_processes_inline(self, ctx):
        result = E.op_ingest_source(ctx, YOUTUBE_URL, run_now=True)
        assert result.data["processed"][0]["status"] == "ready"

    def test_a_spent_balance_blocks_a_new_url(self, ctx, cfg):
        """A URL can only be checked against an empty balance, never a price:
        nothing knows a link's size before fetching it."""
        control_db.debit_credits(ctx.control, ctx.user_id, cfg.quota_credits, cfg=cfg)
        result = E.op_ingest_source(ctx, YOUTUBE_URL)
        assert not result.ok and result.error_reason == "out_of_credits"
        assert f"{cfg.quota_credits:,}" in result.error_message
        assert "first" in result.error_message, "says when they come back"

    def test_an_upload_is_refused_at_the_button_with_its_real_price(self, ctx, cfg, tmp_path):
        """The one case where the price IS knowable before the work: the
        bytes are already on local disk and parsing them costs no LLM call.
        So the person is told the number instead of finding out from a failed
        job minutes later."""
        path = tmp_path / "n.txt"
        path.write_text("some notes that are long enough to chunk into something" * 20)

        from app.ingestion import gatekeeper as gk

        price = E._upload_price(gk.validate_file(path, cfg=cfg), cfg=cfg)
        assert price > 0

        control_db.debit_credits(ctx.control, ctx.user_id, cfg.quota_credits - price + 1, cfg=cfg)
        result = E.op_ingest_file(ctx, path)
        assert not result.ok and result.error_reason == "out_of_credits"
        assert f"{price:,}" in result.error_message, "the real number, not a guess"

    def test_an_image_has_no_price_until_it_is_read(self, ctx, cfg, tmp_path):
        """Honest `None`, not a guessed flat rate. The vision call IS the
        fetch, so an image has no characters until the money is spent — and
        the overdraw is bounded by one response."""
        path = tmp_path / "shot.png"
        path.write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 256)

        from app.ingestion import gatekeeper as gk

        assert E._upload_price(gk.validate_file(path, cfg=cfg), cfg=cfg) is None

    def test_a_file_over_the_precount_ceiling_has_no_price_either(self, ctx, cfg, tmp_path):
        path = tmp_path / "big.txt"
        path.write_text("x" * (cfg.credit_precount_max_bytes + 1))

        from app.ingestion import gatekeeper as gk

        assert E._upload_price(gk.validate_file(path, cfg=cfg), cfg=cfg) is None


class TestYouTubeRefusalIsDiagnosable:
    """Three of the most public videos on YouTube failed in production with
    "private, members-only or age-restricted", and the raw yt-dlp output was
    thrown away so nobody could see why.

    The old test was `any(k in stderr for k in ("private", "members-only",
    "sign in", "age"))`. Both halves of that were wrong.
    """

    def _reason(self, stderr: str) -> str:
        from app.ingestion.fetchers import _youtube_refusal

        return _youtube_refusal(stderr, "https://youtu.be/x").reason

    def test_the_bot_wall_is_not_reported_as_a_private_video(self):
        """YouTube blocking the caller is not something the owner can fix by
        changing the video — telling them to make it public is a dead end."""
        stderr = "ERROR: [youtube] x: Sign in to confirm you're not a bot."
        assert self._reason(stderr) == "platform_blocked_caller"

    def test_a_genuinely_private_video_still_says_so(self):
        assert self._reason("ERROR: [youtube] x: Private video") == "blocked_by_platform"
        assert self._reason("ERROR: members-only content") == "blocked_by_platform"
        assert self._reason("ERROR: age-restricted, sign in") == "blocked_by_platform"

    def test_the_word_message_is_not_an_age_restriction(self):
        """`"age"` was a bare substring, so "message", "usage" and "package"
        all read as age restrictions."""
        for noise in ("Unexpected message from the server",
                      "Exceeded usage quota",
                      "Missing package ffmpeg"):
            assert self._reason(f"ERROR: {noise}") == "fetch_failed", noise

    def test_every_reason_it_can_raise_has_a_remedy(self):
        import re
        from pathlib import Path

        root = Path(__file__).resolve().parents[1]
        frames = (root / "app/static/frames.js").read_text()
        block = frames[frames.index("const REMEDY = {") : frames.index('/** §7.7: "Never')]
        mapped = set(re.findall(r"^  ([a-z_]+):", block, re.M))
        for reason in ("platform_blocked_caller", "blocked_by_platform", "fetch_failed"):
            assert reason in mapped, reason


class TestAMissingJsRuntimeBlamesTheServer:
    """Production's real failure, and my own classifier got it wrong first.

    yt-dlp needs a JavaScript runtime to read YouTube. The image had none, so
    every video failed with "No title found in player response" — and the
    first version of `_youtube_refusal` called that a bot wall, which blames
    YouTube for a dependency we own and sends the operator looking outward.
    """

    def _reason(self, stderr: str) -> str:
        from app.ingestion.fetchers import _youtube_refusal

        return _youtube_refusal(stderr, "https://youtu.be/x").reason

    def test_the_real_production_stderr_is_named_correctly(self):
        stderr = (
            "WARNING: [youtube] No supported JavaScript runtime could be found. Only deno "
            "is enabled by default; to use another runtime add --js-runtimes RUNTIME[:PATH]. "
            "YouTube extraction without a JS runtime has been deprecated. "
            "WARNING: [youtube] No title found in player response"
        )
        assert self._reason(stderr) == "js_runtime_missing"

    def test_ours_is_checked_before_theirs(self):
        """A stderr carrying both must blame the dependency we control."""
        stderr = (
            "No supported JavaScript runtime could be found. "
            "ERROR: Sign in to confirm you're not a bot."
        )
        assert self._reason(stderr) == "js_runtime_missing"

    def test_doctor_fails_when_no_runtime_is_installed(self, cfg, monkeypatch, capsys):
        """The check lives in the prod block, because that is the image that
        has to carry the runtime."""
        import shutil as sh

        from app import cli
        from app import config as config_module

        monkeypatch.setenv("BRAIN_ENV", "prod")
        monkeypatch.setenv("BRAIN_PUBLIC_URL", "https://example.test")
        monkeypatch.setenv("BRAIN_JWT_SECRET", "doctor-secret-key-long-enough-for-hs256")
        config_module.reload_settings()

        real = sh.which
        monkeypatch.setattr(
            cli.shutil, "which",
            lambda name: "" if name in ("deno", "node", "bun", "qjs") else real(name),
        )
        cli.cmd_doctor(cli.build_parser().parse_args(["doctor"]))
        out = capsys.readouterr().out
        assert "JavaScript runtime" in out
        assert "0 failure(s)" not in out, "doctor called itself ready with no JS runtime"

    def test_doctor_is_happy_when_a_runtime_is_there(self, cfg, monkeypatch, capsys):
        import shutil as sh

        from app import cli
        from app import config as config_module

        monkeypatch.setenv("BRAIN_ENV", "prod")
        monkeypatch.setenv("BRAIN_PUBLIC_URL", "https://example.test")
        monkeypatch.setenv("BRAIN_JWT_SECRET", "doctor-secret-key-long-enough-for-hs256")
        config_module.reload_settings()

        real = sh.which
        monkeypatch.setattr(
            cli.shutil, "which",
            lambda name: "/usr/local/bin/deno" if name == "deno" else real(name),
        )
        cli.cmd_doctor(cli.build_parser().parse_args(["doctor"]))
        out = capsys.readouterr().out
        assert "yt-dlp JavaScript runtime" in out and "deno" in out
